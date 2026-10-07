#!/usr/bin/env python3
"""
build.py
Adaptive Distributed AI Inference in Heterogeneous Edge Networks
Research Seminar PPTX – Light Theme Rebuild
Design: #F7F5F0 bg, Calibri only, 20pt+ body, 16pt+ diagram, 14pt+ footer.
Real XML connectors. ~30 slides.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR_TYPE
from pptx.oxml.ns import qn
from lxml import etree

# ── Canvas ─────────────────────────────────────────────────────────────────
W, H  = Inches(13.333), Inches(7.5)
MAR   = Inches(0.6)
CT    = Inches(1.72)          # content top
CB    = Inches(6.8)           # content bottom
CW    = W - 2 * MAR          # ~12.13 in
CH    = CB - CT               # ~5.08 in

# ── Palette ────────────────────────────────────────────────────────────────
BG   = RGBColor(0xF7, 0xF5, 0xF0)   # light bg – all content slides
BGD  = RGBColor(0x1A, 0x1A, 0x1A)   # charcoal – section dividers
INK  = RGBColor(0x1A, 0x1A, 0x1A)   # main text
GRY  = RGBColor(0x6B, 0x6B, 0x6B)   # secondary text / footer
SRF  = RGBColor(0xEC, 0xEA, 0xE4)   # light surface / card bg
BDR  = RGBColor(0xCF, 0xCB, 0xC2)   # border
WHT  = RGBColor(0xFF, 0xFF, 0xFF)
ABL  = RGBColor(0x1F, 0x4E, 0x8C)   # ASET blue
FTL  = RGBColor(0x0F, 0x6B, 0x6B)   # Flocky teal
DOR  = RGBColor(0xC9, 0x77, 0x2E)   # DATOR orange
RED  = RGBColor(0xB0, 0x4A, 0x4A)   # failure red

# ── Font sizes (NOTHING below 14pt) ────────────────────────────────────────
FT  = Pt(34)   # slide title
FS  = Pt(14)   # section label, footer, captions
FB  = Pt(20)   # body / bullets  (min for body text)
FC  = Pt(22)   # card headings
FD  = Pt(16)   # diagram labels (min for diagrams)
FM  = Pt(72)   # big metric numbers

# ── Presentation ──────────────────────────────────────────────────────────
prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
_blank = prs.slide_layouts[6]

# ══════════════════════════════════════════════════════════════════════════
# LOW-LEVEL PRIMITIVES
# ══════════════════════════════════════════════════════════════════════════

def _to_ellipse(shape):
    """Convert a rectangle shape to an ellipse via XML."""
    sp   = shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None: return
    pg = spPr.find(qn('a:prstGeom'))
    if pg is None: return
    pg.set('prst', 'ellipse')
    avLst = pg.find(qn('a:avLst'))
    if avLst is not None:
        for ch in list(avLst): avLst.remove(ch)

def _to_roundrect(shape, pct=8):
    """Convert rectangle to rounded rectangle. pct = corner%."""
    sp   = shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None: return
    pg = spPr.find(qn('a:prstGeom'))
    if pg is None: return
    pg.set('prst', 'roundRect')
    avLst = pg.find(qn('a:avLst'))
    if avLst is None:
        avLst = etree.SubElement(pg, qn('a:avLst'))
    for ch in list(avLst): avLst.remove(ch)
    gd = etree.SubElement(avLst, qn('a:gd'))
    gd.set('name', 'adj')
    gd.set('fmla', f'val {pct * 1000}')

def _add_tailend(conn_shape):
    """Add triangle arrowhead at the end (tail) of a connector."""
    sp   = conn_shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None: return
    ln = spPr.find(qn('a:ln'))
    if ln is None:
        ln = etree.SubElement(spPr, qn('a:ln'))
    for te in ln.findall(qn('a:tailEnd')): ln.remove(te)
    tail = etree.SubElement(ln, qn('a:tailEnd'))
    tail.set('type', 'triangle')
    tail.set('w',   'med')
    tail.set('len', 'med')

def _set_dash(conn_shape):
    """Set connector to dashed line."""
    sp   = conn_shape._element
    spPr = sp.find(qn('p:spPr'))
    if spPr is None: return
    ln = spPr.find(qn('a:ln'))
    if ln is None: ln = etree.SubElement(spPr, qn('a:ln'))
    pd = ln.find(qn('a:prstDash'))
    if pd is None: pd = etree.SubElement(ln, qn('a:prstDash'))
    pd.set('val', 'dash')

def _no_autofit(tf):
    """Disable auto-fit on a TextFrame."""
    tb = tf._txBody
    bp = tb.find(qn('a:bodyPr'))
    if bp is not None:
        bp.set('autofit', 'noAutofit')
        bp.attrib.pop('spAutoFit', None)

# ══════════════════════════════════════════════════════════════════════════
# DRAWING PRIMITIVES
# ══════════════════════════════════════════════════════════════════════════

def R(slide, x, y, w, h, fill=SRF, border=BDR, bw=1.0, round_pct=0):
    """Rectangle (optionally rounded)."""
    sh = slide.shapes.add_shape(1, int(x), int(y), int(w), int(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border; sh.line.width = Pt(bw)
    if round_pct > 0: _to_roundrect(sh, round_pct)
    return sh

def C(slide, cx, cy, d=Inches(0.5), fill=SRF, border=BDR, bw=1.5):
    """Circle centred at (cx, cy) with diameter d."""
    sh = slide.shapes.add_shape(1, int(cx - d//2), int(cy - d//2),
                                 int(d), int(d))
    _to_ellipse(sh)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border; sh.line.width = Pt(bw)
    return sh

def Conn(slide, x1, y1, x2, y2, color=BDR, width=1.0, dashed=False):
    """Plain connector line."""
    c = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed: _set_dash(c)
    return c

def Arrow(slide, x1, y1, x2, y2, color=INK, width=2.0, dashed=False):
    """Connector with arrowhead at (x2, y2)."""
    c = slide.shapes.add_connector(
        MSO_CONNECTOR_TYPE.STRAIGHT,
        int(x1), int(y1), int(x2), int(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed: _set_dash(c)
    _add_tailend(c)
    return c

# ══════════════════════════════════════════════════════════════════════════
# TEXT HELPERS
# ══════════════════════════════════════════════════════════════════════════

def T(slide, text, x, y, w, h,
      sz=FB, bold=False, color=INK,
      align=PP_ALIGN.LEFT, wrap=True, italic=False):
    """Single-paragraph textbox. Explicit font size on every run."""
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    _no_autofit(tf)
    tf.margin_left   = Pt(0)
    tf.margin_right  = Pt(0)
    tf.margin_top    = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name   = 'Calibri'
    r.font.size   = sz
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return tb

def TM(slide, lines, x, y, w, h,
       sz=FB, bold=False, color=INK,
       align=PP_ALIGN.LEFT, wrap=True):
    """
    Multi-paragraph textbox.
    lines: list of str | (text, sz, bold, color)
    """
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    _no_autofit(tf)
    tf.margin_left   = Pt(0)
    tf.margin_right  = Pt(0)
    tf.margin_top    = Pt(2)
    tf.margin_bottom = Pt(2)
    first = True
    for item in lines:
        if isinstance(item, str): t, s, b, c = item, sz, bold, color
        else:
            t = item[0]
            s = item[1] if len(item) > 1 else sz
            b = item[2] if len(item) > 2 else bold
            c = item[3] if len(item) > 3 else color
        if first: p = tf.paragraphs[0]; first = False
        else:     p = tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t; r.font.name = 'Calibri'
        r.font.size = s; r.font.bold = b
        r.font.color.rgb = c
    return tb

def SN(slide, text):
    """Speaker note."""
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    p  = tf.paragraphs[0] if not tf.paragraphs[0].text else tf.add_paragraph()
    r  = p.add_run()
    r.text = text

# ══════════════════════════════════════════════════════════════════════════
# COMPONENT LIBRARY
# ══════════════════════════════════════════════════════════════════════════

def add_bg(slide, color=BG):
    """Full-slide background."""
    R(slide, 0, 0, W, H, fill=color, border=None)

def add_title(slide, section, title, accent=None):
    """
    Section label at (MAR, 0.38), accent bar, title at (MAR, 0.76),
    thin rule at ~1.62 in. Consistent across all content slides.
    """
    sec_col = accent if accent else GRY
    T(slide, section.upper(), MAR, Inches(0.38), CW, Inches(0.35),
      sz=FS, bold=True, color=sec_col)
    if accent:
        R(slide, MAR, Inches(0.70), Inches(2.0), Pt(2), fill=accent, border=None)
    T(slide, title, MAR, Inches(0.76), CW, Inches(0.88),
      sz=FT, bold=True, color=INK)
    R(slide, MAR, Inches(1.62), CW, Pt(1.5), fill=BDR, border=None)

def card(slide, x, y, w, h, heading, body=None, style='light', accent=None):
    """
    Rounded card with optional left accent bar.
    style: 'light' (SRF bg) | 'dark' (BGD bg) | 'neutral'
    body: str or list of str (each 20pt)
    """
    fill_map  = {'light': SRF, 'dark': BGD, 'neutral': BG}
    fill_col  = fill_map.get(style, SRF)
    head_col  = accent if accent else INK
    body_col  = GRY if style != 'dark' else BDR
    text_dark = style != 'dark'

    BAR = Inches(0.07)
    PAD = Inches(0.2)

    R(slide, x, y, w, h, fill=fill_col, border=BDR, bw=1.0, round_pct=8)
    if accent:
        R(slide, x, y, BAR, h, fill=accent, border=None, round_pct=8)

    tx = x + (BAR + PAD if accent else PAD)
    tw = x + w - PAD - tx   # right-side padding

    T(slide, heading, tx, y + PAD, tw, Inches(0.42),
      sz=FC, bold=True, color=head_col if accent else (INK if text_dark else WHT))

    if body:
        if isinstance(body, str): body = [body]
        items = [(b, FB, False, body_col) for b in body]
        by_ = y + PAD + Inches(0.5)
        bh_ = max(h - 2*PAD - Inches(0.5), Inches(0.3))
        TM(slide, items, tx, by_, tw, bh_)

def node(slide, x, y, w, h, title, sub=None, state='normal', accent=None):
    """
    Rounded process/stage box.
    state: 'normal' | 'active' | 'failed' | 'dim'
    """
    fmap = {
        'normal': (SRF, BDR, INK),
        'active': (accent or ABL, accent or ABL, WHT),
        'failed': (RED, RED, WHT),
        'dim':    (BG,  BDR, GRY),
    }
    fill, bdr, tc = fmap.get(state, fmap['normal'])
    R(slide, x, y, w, h, fill=fill, border=bdr, bw=1.5, round_pct=8)

    PAD = Inches(0.12)
    if sub:
        half = (h - 2*PAD) * 0.55
        T(slide, title, x+PAD, y+PAD, w-2*PAD, int(half),
          sz=FD, bold=True, color=tc, align=PP_ALIGN.CENTER, wrap=True)
        T(slide, sub, x+PAD, y+PAD+int(half), w-2*PAD, int((h-2*PAD)*0.4),
          sz=FS, bold=False, color=tc, align=PP_ALIGN.CENTER, wrap=True)
    else:
        T(slide, title, x+PAD, y+PAD, w-2*PAD, h-2*PAD,
          sz=FD, bold=True, color=tc, align=PP_ALIGN.CENTER, wrap=True)

def net_node(slide, cx, cy, state='normal', accent=None, label=None,
             d=Inches(0.55)):
    """
    Circle node for topology diagrams.
    state: 'normal' | 'active' | 'failed' | 'dim'
    """
    fmap = {
        'normal': (SRF, BDR),
        'active': (accent or ABL, accent or ABL),
        'failed': (RED, RED),
        'dim':    (BG,  BDR),
    }
    fill, bdr = fmap.get(state, fmap['normal'])
    sh = C(slide, cx, cy, d=d, fill=fill, border=bdr, bw=2.0)
    if label:
        tc = WHT if state in ('active', 'failed') else INK
        T(slide, label, int(cx - d//2), int(cy - d*0.35),
          int(d), int(d*0.7), sz=FS, bold=True, color=tc,
          align=PP_ALIGN.CENTER, wrap=False)
    return sh

def link(slide, sha, shb, active=False, accent=None):
    """Thin line between shape centres."""
    ax = sha.left + sha.width  // 2
    ay = sha.top  + sha.height // 2
    bx = shb.left + shb.width  // 2
    by = shb.top  + shb.height // 2
    col = (accent or ABL) if active else BDR
    wid = 2.0 if active else 1.0
    return Conn(slide, ax, ay, bx, by, color=col, width=wid)

def arrow(slide, sha, shb, kind='solid', accent=None):
    """
    Real connector with triangle arrowhead from sha edge to shb edge.
    kind: 'solid' | 'dashed' | 'red'
    """
    ax_c = sha.left + sha.width  // 2
    ay_c = sha.top  + sha.height // 2
    bx_c = shb.left + shb.width  // 2
    by_c = shb.top  + shb.height // 2

    if abs(by_c - ay_c) >= abs(bx_c - ax_c):   # primarily vertical
        x1 = sha.left + sha.width // 2
        y1 = (sha.top + sha.height) if by_c > ay_c else sha.top
        x2 = shb.left + shb.width  // 2
        y2 = shb.top if by_c > ay_c else (shb.top + shb.height)
    else:                                         # primarily horizontal
        x1 = (sha.left + sha.width) if bx_c > ax_c else sha.left
        y1 = sha.top + sha.height // 2
        x2 = shb.left if bx_c > ax_c else (shb.left + shb.width)
        y2 = shb.top  + shb.height // 2

    col = accent or INK
    if kind == 'red': col = RED
    return Arrow(slide, x1, y1, x2, y2,
                 color=col, width=2.0, dashed=(kind=='dashed'))

def metric(slide, x, y, w, number, label, source=None, accent=INK):
    """Big number + descriptor label + italic source."""
    T(slide, number, x, y, w, Inches(1.5),
      sz=FM, bold=True, color=accent, align=PP_ALIGN.CENTER)
    T(slide, label, x, y + Inches(1.52), w, Inches(0.55),
      sz=FB, bold=False, color=INK, align=PP_ALIGN.CENTER, wrap=True)
    if source:
        T(slide, source, x, y + Inches(2.15), w, Inches(0.35),
          sz=FS, bold=False, color=GRY, align=PP_ALIGN.CENTER, italic=True)

def add_footer(slide, n, source=None):
    """Slide number (bottom-right) + optional source (bottom-left)."""
    T(slide, str(n), W - MAR - Inches(0.65), H - Inches(0.45),
      Inches(0.65), Inches(0.35), sz=FS, bold=False, color=GRY,
      align=PP_ALIGN.RIGHT)
    if source:
        T(slide, source, MAR, H - Inches(0.45), Inches(8), Inches(0.35),
          sz=FS, bold=False, color=GRY, italic=True)

def section_divider(prs, num_str, name, tagline, accent):
    """
    Charcoal background divider slide.
    Large numeral + system name + tagline + node/line motif.
    """
    sl = prs.slides.add_slide(_blank)
    add_bg(sl, BGD)

    # Large numeral (top-left area)
    T(sl, num_str, MAR, Inches(0.5), Inches(4), Inches(2.5),
      sz=Pt(110), bold=True, color=accent)

    # System name
    T(sl, name, MAR, Inches(3.0), Inches(8), Inches(1.3),
      sz=Pt(54), bold=True, color=WHT)

    # Tagline
    T(sl, tagline, MAR, Inches(4.35), Inches(9), Inches(0.7),
      sz=Pt(22), bold=False, color=GRY)

    # Node/line motif (right side) – decorative
    motif_cx = [Inches(10.3), Inches(11.8), Inches(12.5),
                Inches(11.0), Inches(9.6),  Inches(12.3)]
    motif_cy = [Inches(1.8),  Inches(1.2),  Inches(2.6),
                Inches(3.4),  Inches(3.0),  Inches(4.5)]
    edges = [(0,1),(1,2),(2,3),(3,4),(4,0),(2,5),(3,5)]
    for i, (ax, ay) in enumerate(zip(motif_cx, motif_cy)):
        C(sl, ax, ay, d=Inches(0.4), fill=None, border=accent, bw=1.5)
    for i, j in edges:
        col = accent if (i + j) % 3 != 1 else GRY
        Conn(sl, motif_cx[i], motif_cy[i],
             motif_cx[j], motif_cy[j], color=col, width=0.75)

    return sl

# ══════════════════════════════════════════════════════════════════════════
# SLIDE BUILDERS
# ══════════════════════════════════════════════════════════════════════════

n = 0   # slide counter

# ── Slide 1: Title ─────────────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
# Top rule
R(sl, MAR, Inches(0.55), CW, Pt(2), fill=INK, border=None)
# Main title
T(sl, "ADAPTIVE DISTRIBUTED AI INFERENCE",
  MAR, Inches(0.75), CW, Inches(1.05),
  sz=Pt(38), bold=True, color=INK)
T(sl, "IN HETEROGENEOUS EDGE NETWORKS",
  MAR, Inches(1.72), CW, Inches(0.85),
  sz=Pt(36), bold=True, color=ABL)
# Subtitle
T(sl, "Three research directions  ·  One problem space  ·  From scheduling to orchestration",
  MAR, Inches(2.62), CW, Inches(0.55),
  sz=FB, bold=False, color=GRY)
# Three paper tags
tags = [
    ("Paper 1", "ASET", "Adaptive Scheduling", ABL),
    ("Paper 2", "Flocky", "Decentralized Orchestration", FTL),
    ("Paper 3", "DATOR", "Dynamic Distributed Inference", DOR),
]
tw2 = Inches(3.8)
tx2 = MAR
for ptag, pname, ptheme, pcol in tags:
    R(sl, tx2, Inches(3.4), tw2, Inches(1.7), fill=SRF, border=BDR, bw=1, round_pct=8)
    R(sl, tx2, Inches(3.4), tw2, Inches(0.06), fill=pcol, border=None)
    T(sl, ptag,  tx2 + Inches(0.15), Inches(3.52), tw2, Inches(0.28),
      sz=FS, bold=True, color=pcol)
    T(sl, pname, tx2 + Inches(0.15), Inches(3.78), tw2, Inches(0.55),
      sz=Pt(26), bold=True, color=INK)
    T(sl, ptheme, tx2 + Inches(0.15), Inches(4.32), tw2, Inches(0.55),
      sz=FB, bold=False, color=GRY)
    tx2 += tw2 + Inches(0.27)
# Bottom rule + presenter
R(sl, MAR, Inches(5.3), CW, Pt(1.5), fill=BDR, border=None)
T(sl, "Presented by  Arjun A S  ·  SCT23CS017",
  MAR, Inches(5.5), Inches(8), Inches(0.5),
  sz=FB, bold=True, color=INK)
T(sl, "Dept. of Computer Science & Engineering  ·  Sree Chitra Thirunal College of Engineering",
  MAR, Inches(5.95), CW, Inches(0.42),
  sz=FS, bold=False, color=GRY)
T(sl, "Guide: Prof. Arun K",
  W - MAR - Inches(3.5), Inches(5.95), Inches(3.5), Inches(0.42),
  sz=FS, bold=True, color=GRY, align=PP_ALIGN.RIGHT)
add_footer(sl, n)

# ── Slide 2: Introduction ──────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Motivation", "Why Does Edge AI Need Smarter Coordination?")
# Left: challenge items as cards (3 items)
chy = CT
chw = Inches(5.6); chh = Inches(1.55); chgap = Inches(0.12)
challenges = [
    ("Variable Compute & Load",  ["Nodes: microcontrollers to GPUs", "Hot-spots form; edge saturates"]),
    ("Dynamic Network Latency",  ["Delay shifts across edge hops", "Routes change continuously"]),
    ("Real-Time Inference Needs",["Strict end-to-end latency targets", "Node failures break pipelines"]),
]
for (ch_head, ch_body) in challenges:
    card(sl, MAR, chy, chw, chh, ch_head, ch_body, style='light')
    chy += chh + chgap
# Right: central question box
qx = MAR + chw + Inches(0.5)
qw = CW - chw - Inches(0.5)
R(sl, qx, CT, qw, CH, fill=SRF, border=ABL, bw=2.0, round_pct=8)
T(sl, "THE CENTRAL QUESTION",
  qx + Inches(0.25), CT + Inches(0.25), qw - Inches(0.5), Inches(0.35),
  sz=FS, bold=True, color=ABL)
T(sl, '"Who processes\nthis workload\nright now?"',
  qx + Inches(0.25), CT + Inches(0.7), qw - Inches(0.5), Inches(2.4),
  sz=Pt(28), bold=True, color=INK)
T(sl, "As load, latency, and node state change —\nthe answer changes too.",
  qx + Inches(0.25), CT + Inches(3.3), qw - Inches(0.5), Inches(0.85),
  sz=FB, bold=False, color=GRY)
add_footer(sl, n)

# ── Slide 3: Key Concepts ──────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Foundations", "Core Concepts in Edge AI Coordination")
# 4 cards in 2×2 grid
concepts = [
    ("Heterogeneity", ["Nodes vary: CPU, GPU, OS, memory", "No uniform scheduling rule works"], None),
    ("Scheduling",    ["Select: model variant + resource cluster", "Meet delay, accuracy, throughput"], None),
    ("Orchestration", ["Coordinate deployment across nodes", "Lifecycle and placement management"], None),
    ("Adaptation",    ["Revise decisions as conditions change", "Workload, topology, availability"], None),
]
cw3 = (CW - Inches(0.2)) / 2
ch3 = (CH - Inches(0.15)) / 2
positions = [
    (MAR, CT),
    (MAR + cw3 + Inches(0.2), CT),
    (MAR, CT + ch3 + Inches(0.15)),
    (MAR + cw3 + Inches(0.2), CT + ch3 + Inches(0.15)),
]
for (cx_, cy_), (ch_h, ch_b, acc) in zip(positions, concepts):
    card(sl, cx_, cy_, cw3, ch3, ch_h, ch_b, style='light')
add_footer(sl, n)

# ── Slide 4: Roadmap ───────────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Seminar Roadmap", "Three Papers · Three Lenses on One Problem")
# 3 tall cards
rcard_w = (CW - Inches(0.4)) / 3
rcard_h = CH
rcards = [
    ("ASET", ABL,
     ['"Which scheduling policy', 'should I use right now?"'],
     ["Adaptive scheduling via DRL", "Policy selection meta-scheduler"]),
    ("Flocky", FTL,
     ['"How can nodes coordinate', 'without a central controller?"'],
     ["Intent-based orchestration", "Decentralized discovery + OAM"]),
    ("DATOR", DOR,
     ['"How can inference survive', 'node and link failures?"'],
     ["Self-organized multi-stage", "Dynamic rerouting + resilience"]),
]
rx3 = MAR
for (sysname, acc, q_lines, meth_lines) in rcards:
    R(sl, rx3, CT, rcard_w, rcard_h, fill=SRF, border=BDR, bw=1, round_pct=8)
    R(sl, rx3, CT, rcard_w, Inches(0.06), fill=acc, border=None, round_pct=8)
    T(sl, sysname, rx3 + Inches(0.2), CT + Inches(0.15), rcard_w - Inches(0.4), Inches(0.75),
      sz=Pt(36), bold=True, color=acc)
    qy = CT + Inches(0.95)
    for ql in q_lines:
        T(sl, ql, rx3 + Inches(0.2), qy, rcard_w - Inches(0.4), Inches(0.42),
          sz=FB, bold=False, color=INK)
        qy += Inches(0.42)
    R(sl, rx3 + Inches(0.2), qy + Inches(0.1), rcard_w - Inches(0.4), Pt(1),
      fill=BDR, border=None)
    for ml in meth_lines:
        qy += Inches(0.4)
        T(sl, ml, rx3 + Inches(0.2), qy + Inches(0.12), rcard_w - Inches(0.4), Inches(0.38),
          sz=FB, bold=False, color=GRY)
    rx3 += rcard_w + Inches(0.2)
# Progression label
T(sl, "Scheduling  →  Orchestration  →  Resilient Inference",
  MAR, CB + Inches(0.05), CW, Inches(0.35),
  sz=FS, bold=True, color=GRY, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 5: ASET Section Divider ──────────────────────────────────────────
sl = section_divider(prs, "01", "ASET",
                     "Adaptive Scheduling for Edge & Cloud Inference", ABL); n += 1
add_footer(sl, n)

# ── Slide 6: ASET Paper Info ───────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Paper 1 of 3  ·  ASET",
          "Scheduling Inference Workloads in the Computing Continuum", ABL)
# Left: metadata card
card(sl, MAR, CT, Inches(7.5), CH,
     "Paper Metadata",
     ["Authors: G. Castellano, J-J. Nieto, A. Angi,",
      "   F. Á. Terribas, J. Luque, F. D. Andilla,",
      "   A. Sacco, F. Esposito, F. Risso",
      "Venue: IEEE ICDCS Workshops 2025",
      "Pages: 297–302",
      "DOI: 10.1109/ICDCSW63273.2025.00056"],
     style='light', accent=ABL)
# Right: role card
rx4 = MAR + Inches(7.7)
rw4 = CW - Inches(7.7)
R(sl, rx4, CT, rw4, CH, fill=ABL, border=None, round_pct=8)
T(sl, "ROLE IN\nSEMINAR",
  rx4 + Inches(0.2), CT + Inches(0.2), rw4 - Inches(0.4), Inches(0.9),
  sz=FS, bold=True, color=WHT)
T(sl, "LAYER 1",
  rx4 + Inches(0.2), CT + Inches(1.2), rw4 - Inches(0.4), Inches(0.5),
  sz=FC, bold=True, color=WHT)
T(sl, "Adaptive\nDecision\nMaking",
  rx4 + Inches(0.2), CT + Inches(1.75), rw4 - Inches(0.4), Inches(2.0),
  sz=Pt(26), bold=True, color=WHT)
add_footer(sl, n)
SN(sl, "ASET: reinforcement learning-based adaptive scheduler for DNN inference across edge-cloud continuum.")

# ── Slide 7: ASET Problem ──────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "ASET  ·  Problem Formulation",
          "Scheduling DNN Inference Across the Continuum", ABL)
# 3 column inputs
col_data = [
    ("Stream Requirements", ["Max end-to-end delay", "Min accuracy (mAP)", "Input / data rate", "Input size"]),
    ("Model Variants",      ["Accuracy · delay · throughput", "CPU/GPU · batch size", "Number of replicas", "Deployment config"]),
    ("Resource Clusters",   ["Edge / access network", "Central-office resources", "ISP / data-centre", "Cloud resources"]),
]
col_w4 = (CW - Inches(0.4)) / 3
col_h4 = Inches(3.2)
col_x  = MAR
for (chead, citems) in col_data:
    card(sl, col_x, CT, col_w4, col_h4, chead, citems, style='light', accent=ABL)
    col_x += col_w4 + Inches(0.2)
# Arrow down: using coordinatesh
mid_x = MAR + CW // 2
arrow_y1 = CT + col_h4
arrow_y2 = arrow_y1 + Inches(0.55)
Arrow(sl, mid_x, arrow_y1, mid_x, arrow_y2, color=ABL, width=2.5)
# Decision box
dec_y = arrow_y2 + Inches(0.02)
dec_h = Inches(0.9)
R(sl, MAR, dec_y, CW, dec_h, fill=ABL, border=None, round_pct=6)
T(sl, "Select: feasible model variant  +  resource cluster  →  all constraints satisfied",
  MAR + Inches(0.3), dec_y + Inches(0.08), CW - Inches(0.6), dec_h - Inches(0.16),
  sz=FC, bold=True, color=WHT, align=PP_ALIGN.CENTER)
# Objective
obj_y = dec_y + dec_h + Inches(0.2)
R(sl, MAR, obj_y, CW, Inches(0.55), fill=SRF, border=ABL, bw=1.5, round_pct=6)
T(sl, "Objective: maximise the fraction of successfully served inference queries",
  MAR + Inches(0.3), obj_y + Inches(0.07), CW - Inches(0.6), Inches(0.38),
  sz=FB, bold=False, color=INK, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 8: ASET Static Policies ─────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "ASET  ·  Static Scheduling Policies",
          "Why One Fixed Policy Is Insufficient", ABL)
policies = [
    ("Closest",          "Minimise network distance / delay"),
    ("Load Balance",     "Prefer resource with lowest current load"),
    ("Farthest",         "Preserve nearby capacity; use distant nodes"),
    ("Cheaper",          "Cost-oriented within available delay budget"),
    ("Least Impedance",  "Minimise end-to-end latency"),
]
pw5 = Inches(5.7)
ph5 = Inches(0.92)
pgap = Inches(0.11)
py5 = CT
for (pname, pdesc) in policies:
    card(sl, MAR, py5, pw5, ph5, pname, [pdesc], style='light')
    py5 += ph5 + pgap
# Right: insight box
rx5 = MAR + pw5 + Inches(0.35)
rw5 = CW - pw5 - Inches(0.35)
R(sl, rx5, CT, rw5, CH, fill=SRF, border=RED, bw=2, round_pct=8)
T(sl, "THE LIMITATION", rx5 + Inches(0.2), CT + Inches(0.2),
  rw5 - Inches(0.4), Inches(0.35), sz=FS, bold=True, color=RED)
insights = [
    ("Low load →", "nearby edge is ideal"),
    ("High load →", "same node overloads"),
    ("Changing rates →", "best policy shifts"),
]
iy = CT + Inches(0.65)
for (lab, val) in insights:
    T(sl, lab, rx5 + Inches(0.2), iy, rw5 - Inches(0.4), Inches(0.38),
      sz=FB, bold=True, color=INK)
    T(sl, val, rx5 + Inches(0.2), iy + Inches(0.38), rw5 - Inches(0.4), Inches(0.38),
      sz=FB, bold=False, color=GRY)
    iy += Inches(0.9)
R(sl, rx5 + Inches(0.2), iy + Inches(0.1), rw5 - Inches(0.4), Pt(1.5),
  fill=BDR, border=None)
T(sl, "No single policy stays\noptimal across all conditions.",
  rx5 + Inches(0.2), iy + Inches(0.25), rw5 - Inches(0.4), Inches(1.1),
  sz=FC, bold=True, color=RED, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 9: ASET Meta-Scheduler (RL Loop) ────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "ASET  ·  RL Methodology",
          "ASET as a Meta-Scheduler", ABL)
# RL loop: 6 boxes connected by arrows
loop_items = [
    ("System State",       "Apps · network · load · arrivals",    'normal'),
    ("RL Agent (DQN)",     "Estimates Q-values for each policy",  'active'),
    ("Select Policy",      "e.g. Load Balance / Closest / …",    'normal'),
    ("Policy Executes",    "Maps streams → variants → clusters",  'normal'),
    ("Outcome Observed",   "QoS success / service quality",       'normal'),
    ("Reward Signal",      "Successful query ratio",              'normal'),
]
lbox_w = Inches(5.8)
lbox_h = Inches(0.72)
lbox_x = MAR
lbox_gap = Inches(0.1)
lbox_shapes = []
ly = CT
for (ltitle, lsub, lstate) in loop_items:
    acc = ABL if lstate == 'active' else None
    node(sl, lbox_x, ly, lbox_w, lbox_h, ltitle, sub=lsub,
         state=lstate, accent=ABL)
    lbox_shapes.append(
        type('Sh', (), {'left':lbox_x, 'top':ly,
                        'width':lbox_w, 'height':lbox_h})())
    ly += lbox_h + lbox_gap
# Arrows between boxes
for i in range(len(lbox_shapes) - 1):
    sa = lbox_shapes[i]; sb = lbox_shapes[i+1]
    Arrow(sl, sa.left + sa.width//2, sa.top + sa.height,
          sb.left + sb.width//2, sb.top, color=ABL, width=2.0)
# Feedback loop curve (right side)
fb_x = lbox_x + lbox_w + Inches(0.25)
fb_top = CT + lbox_h // 2
fb_bot = CT + 5 * (lbox_h + lbox_gap) + lbox_h // 2
R(sl, fb_x, fb_top, Inches(0.03), fb_bot - fb_top, fill=ABL, border=None)
R(sl, lbox_x + lbox_w, fb_top, Inches(0.28), Inches(0.03), fill=ABL, border=None)
R(sl, lbox_x + lbox_w, fb_bot, Inches(0.28), Inches(0.03), fill=ABL, border=None)
T(sl, "Feedback\nLoop", fb_x + Inches(0.05), CT + Inches(2.2), Inches(1.0), Inches(0.65),
  sz=FS, bold=True, color=ABL)
# Right panel
rx6 = MAR + lbox_w + Inches(0.65)
rw6 = CW - lbox_w - Inches(0.65)
R(sl, rx6, CT, rw6, CH, fill=SRF, border=ABL, bw=1.5, round_pct=8)
T(sl, "KEY DISTINCTION", rx6 + Inches(0.2), CT + Inches(0.2),
  rw6 - Inches(0.4), Inches(0.35), sz=FS, bold=True, color=ABL)
T(sl, "ASET does NOT learn\nwhich node to use.",
  rx6 + Inches(0.2), CT + Inches(0.65), rw6 - Inches(0.4), Inches(0.95),
  sz=FC, bold=True, color=INK)
T(sl, "It learns WHICH POLICY\nto apply for the current state.",
  rx6 + Inches(0.2), CT + Inches(1.65), rw6 - Inches(0.4), Inches(0.95),
  sz=FC, bold=True, color=ABL)
T(sl, "T = 25 s scheduling window",
  rx6 + Inches(0.2), CT + Inches(2.8), rw6 - Inches(0.4), Inches(0.45),
  sz=FB, bold=False, color=GRY)
T(sl, "Trained: multiple topologies\nand workload configurations",
  rx6 + Inches(0.2), CT + Inches(3.3), rw6 - Inches(0.4), Inches(0.85),
  sz=FB, bold=False, color=GRY)
add_footer(sl, n)

# ── Slide 10: ASET DQN Architecture ───────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "ASET  ·  Deep Q-Network",
          "MDP Formulation & DQN Architecture", ABL)
# Left: MDP table
mdp_rows = [
    ("State  S_t",   "Apps · network topology · arrivals · load"),
    ("Action  A_t",  "Choice of scheduling policy"),
    ("Reward",       "Successful query servicing quality"),
    ("Transition",   "Next state after policy applied for T=25s"),
]
mw = Inches(6.0); mh_row = Inches(0.98); my_ = CT
for (mlabel, mdesc) in mdp_rows:
    R(sl, MAR, my_, mw, mh_row, fill=SRF, border=BDR, bw=1.0)
    R(sl, MAR, my_, Inches(0.06), mh_row, fill=ABL, border=None)
    T(sl, mlabel, MAR + Inches(0.15), my_ + Inches(0.08),
      Inches(1.8), mh_row - Inches(0.16), sz=FD, bold=True, color=ABL)
    T(sl, mdesc, MAR + Inches(2.1), my_ + Inches(0.08),
      mw - Inches(2.3), mh_row - Inches(0.16), sz=FB, bold=False, color=GRY)
    my_ += mh_row + Inches(0.1)
# Right: DQN stack
dw = CW - mw - Inches(0.5)
dx = MAR + mw + Inches(0.5)
dqn_layers = [
    ("State Input",        "SRF", ABL),
    ("Conv Layers",        "ABL", "WHT"),
    ("Fully Connected",    "SRF", ABL),
    ("Q-values (per policy)", "ABL", "WHT"),
]
dh_box = Inches(0.82)
dgap   = Inches(0.2)
dy_    = CT
dshapes = []
for (dlabel, dfill_key, dtxt_key) in dqn_layers:
    dfill = SRF if dfill_key == "SRF" else ABL
    dtxt  = WHT if dtxt_key == "WHT" else ABL
    node(sl, dx, dy_, dw, dh_box, dlabel, state='active' if dfill_key == 'ABL' else 'normal', accent=ABL)
    dshapes.append(type('S',(),{'left':dx,'top':dy_,'width':dw,'height':dh_box})())
    dy_ += dh_box + dgap
for i in range(len(dshapes)-1):
    sa = dshapes[i]; sb = dshapes[i+1]
    Arrow(sl, sa.left+sa.width//2, sa.top+sa.height,
          sb.left+sb.width//2, sb.top, color=ABL, width=2)
# Q-value note
T(sl, "Q(state, policy) = expected return for selecting that policy now",
  dx, dy_ + Inches(0.08), dw, Inches(0.45),
  sz=FB, bold=False, color=GRY, align=PP_ALIGN.CENTER, wrap=True)
T(sl, "ε-greedy exploration · dynamic epsilon update",
  dx, dy_ + Inches(0.55), dw, Inches(0.38),
  sz=FB, bold=False, color=GRY, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 11: ASET Evaluation ─────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "ASET  ·  Evaluation",
          "Results Across Topologies & Workloads", ABL)
topo_cards = [
    ("dc-cloud",    ABL,
     ["ASET not uniformly best", "Performance is topology-dependent",
      "Can underperform some statics in higher load"]),
    ("co-dc-cloud", ABL,
     ["Improves over evaluated baselines", "Holds even outside training arrival rate",
      "Moderate generalization gains"]),
    ("full-edge",   ABL,
     ["Strongest improvement", "↑ successful-query ratio",
      "Generalises to dynamic workloads"]),
]
tcw = (CW - Inches(0.4)) / 3
tch = Inches(3.2)
tx7 = MAR
for (tname, tacc, tbody) in topo_cards:
    card(sl, tx7, CT, tcw, tch, tname, tbody, style='light', accent=ABL)
    tx7 += tcw + Inches(0.2)
# Takeaway
tak_y = CT + tch + Inches(0.2)
tak_h = CH - tch - Inches(0.2)
R(sl, MAR, tak_y, CW, tak_h, fill=ABL, border=None, round_pct=8)
T(sl, "Key Takeaway",
  MAR + Inches(0.3), tak_y + Inches(0.15), CW - Inches(0.6), Inches(0.38),
  sz=FS, bold=True, color=WHT)
T(sl, "State-aware policy selection adapts across conditions.\nPerformance gains are topology-specific — not universal.",
  MAR + Inches(0.3), tak_y + Inches(0.55), CW - Inches(0.6), tak_h - Inches(0.7),
  sz=FC, bold=True, color=WHT, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 12: Flocky Section Divider ──────────────────────────────────────
sl = section_divider(prs, "02", "Flocky",
                     "Decentralized Intent-based Edge Orchestration", FTL); n += 1
add_footer(sl, n)

# ── Slide 13: Flocky Paper Info ────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Paper 2 of 3  ·  Flocky",
          "Decentralized Intent-based Edge Orchestration Using OAM", FTL)
card(sl, MAR, CT, Inches(7.5), CH,
     "Paper Metadata",
     ["Authors: T. Goethals, M. Sebrechts, M. Al-Naday,",
      "   F. De Turck, B. Volckaert",
      "Venue: IEEE Trans. Services Computing",
      "Volume 19, Issue 1, pp. 380–395, 2026",
      "DOI: 10.1109/TSC.2025.3631387"],
     style='light', accent=FTL)
rx8 = MAR + Inches(7.7)
rw8 = CW - Inches(7.7)
R(sl, rx8, CT, rw8, CH, fill=FTL, border=None, round_pct=8)
T(sl, "ROLE IN\nSEMINAR",
  rx8 + Inches(0.2), CT + Inches(0.2), rw8 - Inches(0.4), Inches(0.9),
  sz=FS, bold=True, color=WHT)
T(sl, "LAYER 2",
  rx8 + Inches(0.2), CT + Inches(1.2), rw8 - Inches(0.4), Inches(0.5),
  sz=FC, bold=True, color=WHT)
T(sl, "Decentralized\nCoordination",
  rx8 + Inches(0.2), CT + Inches(1.75), rw8 - Inches(0.4), Inches(2.0),
  sz=Pt(26), bold=True, color=WHT)
add_footer(sl, n)
SN(sl, "Flocky: intent-based decentralized edge orchestration using the Open Application Model.")

# ── Slide 14: Flocky Problem (Cloud vs Edge) ───────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Flocky  ·  Problem",
          "Cloud Assumptions Don't Hold at the Edge")
# Left: cloud column (neutral)
cld_w = (CW - Inches(0.5)) / 2
card(sl, MAR, CT, cld_w, CH,
     "Cloud Assumptions",
     ["Small set of powerful machines",
      "Centralized orchestrator (Kubernetes)",
      "High-capacity, stable infrastructure",
      "Fast, reliable networking"],
     style='neutral')
# VS divider
T(sl, "VS",
  MAR + cld_w + Inches(0.05), CT + CH//2 - Inches(0.25),
  Inches(0.4), Inches(0.5),
  sz=FC, bold=True, color=BDR, align=PP_ALIGN.CENTER)
# Right: edge column
edg_x = MAR + cld_w + Inches(0.5)
edg_w = CW - cld_w - Inches(0.5)
card(sl, edg_x, CT, edg_w, CH,
     "Edge Reality",
     ["Many heterogeneous, limited nodes",
      "No single entity has global view",
      "Changing topology and connectivity",
      "Resource-constrained devices"],
     style='light', accent=FTL)
# Bottom statement
R(sl, MAR, CB - Inches(0.02), CW, Pt(1.5), fill=FTL, border=None)
T(sl, "Can orchestration itself become decentralized?",
  MAR, CB + Inches(0.05), CW, Inches(0.35),
  sz=FB, bold=True, color=FTL, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 15: Flocky Intent & OAM ─────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Flocky  ·  Methodology",
          "From Workload Intent to Placement Decision", FTL)
# 5-step vertical flow (left column)
flow_items = [
    ("Workload Intent",     "Requirements expressed declaratively"),
    ("Open Application Model", "Structured component description + edge traits"),
    ("Node Metadata",       "Per-node capability advertisement"),
    ("Capability Matching", "Requirements ↔ metadata per candidate node"),
    ("Placement Decision",  "Best-matching node(s) selected & deployed"),
]
fw = Inches(6.2)
fh = Inches(0.82)
fgap = Inches(0.1)
fy_  = CT
fshapes = []
for (fname, fdesc) in flow_items:
    node(sl, MAR, fy_, fw, fh, fname, sub=fdesc, state='normal', accent=FTL)
    fshapes.append(type('S',(),{'left':MAR,'top':fy_,'width':fw,'height':fh})())
    fy_ += fh + fgap
for i in range(len(fshapes)-1):
    sa = fshapes[i]; sb = fshapes[i+1]
    Arrow(sl, sa.left+sa.width//2, sa.top+sa.height,
          sb.left+sb.width//2, sb.top, color=FTL, width=2)
# Right: OAM traits panel
rx9 = MAR + fw + Inches(0.5)
rw9 = CW - fw - Inches(0.5)
R(sl, rx9, CT, rw9, CH, fill=SRF, border=FTL, bw=1.5, round_pct=8)
T(sl, "OAM EDGE TRAITS", rx9 + Inches(0.2), CT + Inches(0.2),
  rw9 - Inches(0.4), Inches(0.35), sz=FS, bold=True, color=FTL)
T(sl, "Orchestration considers\nmore than raw CPU:",
  rx9 + Inches(0.2), CT + Inches(0.65), rw9 - Inches(0.4), Inches(0.7),
  sz=FB, bold=False, color=GRY)
traits = ["Network encryption", "Attestation", "Green energy",
          "Secure runtime", "Distance constraints"]
ty_ = CT + Inches(1.4)
for tr in traits:
    R(sl, rx9 + Inches(0.2), ty_, rw9 - Inches(0.4), Inches(0.52),
      fill=BG, border=FTL, bw=0.75, round_pct=6)
    T(sl, f"  ✦  {tr}", rx9 + Inches(0.3), ty_ + Inches(0.05),
      rw9 - Inches(0.6), Inches(0.42), sz=FB, bold=False, color=INK)
    ty_ += Inches(0.65)
add_footer(sl, n)

# ── Slide 16: Flocky Discovery (real network diagram) ─────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Flocky  ·  Decentralized Discovery",
          "Building Local Topology Knowledge", FTL)
# Real network diagram — 8 nodes, node A selected
ND = Inches(0.6)   # node diameter
# Positions: (cx, cy) in absolute coords
npx = [Inches(3.4), Inches(2.0), Inches(4.8), Inches(1.4),
       Inches(5.5), Inches(2.8), Inches(5.0), Inches(3.5)]
npy = [Inches(3.6), Inches(2.8), Inches(2.5), Inches(4.4),
       Inches(3.0), Inches(5.0), Inches(4.7), Inches(2.1)]
nlabels = ["A", "B", "C", "D", "E", "F", "G", "H"]
nstates  = ['active', 'active', 'active', 'dim', 'normal', 'dim', 'dim', 'normal']
# Links: (i, j, active)
net_links = [
    (0,1,True),(0,2,True),(0,5,False),(0,7,False),
    (1,3,False),(1,5,False),(2,4,True),(2,7,True),
    (3,5,False),(4,6,False),(6,5,False),
]
# Draw links first (below nodes)
node_shapes_f = []
for (li, lj, la) in net_links:
    col = FTL if la else BDR
    wid = 2.0 if la else 0.75
    Conn(sl, npx[li], npy[li], npx[lj], npy[lj], color=col, width=wid)
# Draw nodes
for i in range(len(npx)):
    sh = net_node(sl, npx[i], npy[i], state=nstates[i], accent=FTL,
                  label=nlabels[i], d=ND)
    node_shapes_f.append(sh)
# Explanation text (right side)
rx10 = Inches(7.1)
rw10 = W - MAR - rx10
T(sl, "DISCOVERY STEPS", rx10, CT + Inches(0.0), rw10, Inches(0.38),
  sz=FS, bold=True, color=FTL)
disc_steps = [
    "1.  Local node broadcasts capability info",
    "2.  Nearby nodes respond with metadata",
    "3.  PING measures round-trip latency",
    "4.  Metadata cached for placement decisions",
    "5.  Local topology view built — no global state",
]
dsy = CT + Inches(0.45)
for ds in disc_steps:
    T(sl, ds, rx10, dsy, rw10, Inches(0.52),
      sz=FB, bold=False, color=INK)
    dsy += Inches(0.62)
R(sl, rx10, dsy + Inches(0.1), rw10, Inches(0.9),
  fill=SRF, border=FTL, bw=1.5, round_pct=6)
T(sl, ">97% of neighbours\ndiscovered in 2 rounds",
  rx10 + Inches(0.15), dsy + Inches(0.15), rw10 - Inches(0.3), Inches(0.65),
  sz=FB, bold=True, color=FTL, italic=True)
# Legend
T(sl, "■ Teal = selected node + local links   ■ Grey = distant / unvisited",
  MAR, CB + Inches(0.05), Inches(6.5), Inches(0.32),
  sz=FS, bold=False, color=GRY)
add_footer(sl, n, source="As reported by the paper")

# ── Slide 17: Flocky Orchestration Pipeline ────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Flocky  ·  Orchestration Pipeline",
          "Five-Stage Decentralized Orchestration", FTL)
# 5 stage nodes, horizontal flow
stage_names = ["1\nDiscovery", "2\nMetadata", "3\nScheduling",
               "4\nConstruct", "5\nDeployment"]
stage_desc  = [
    "Find nearby nodes",
    "Learn capabilities",
    "Match requirements",
    "Compose workload",
    "Place and launch",
]
stage_w = (CW - Inches(0.8)) / 5
stage_h = Inches(2.8)
stage_x = MAR
sshapes = []
for (sn_txt, sd_txt) in zip(stage_names, stage_desc):
    node(sl, stage_x, CT, stage_w, stage_h, sn_txt, sub=sd_txt,
         state='active', accent=FTL)
    sshapes.append(type('S',(),{
        'left': stage_x, 'top': CT,
        'width': stage_w, 'height': stage_h
    })())
    stage_x += stage_w + Inches(0.2)
for i in range(len(sshapes)-1):
    sa = sshapes[i]; sb = sshapes[i+1]
    Arrow(sl, sa.left+sa.width, sa.top+sa.height//2,
          sb.left, sb.top+sb.height//2, color=FTL, width=2.5)
# Distinction box
dist_y = CT + stage_h + Inches(0.3)
R(sl, MAR, dist_y, CW, CH - stage_h - Inches(0.3), fill=SRF, border=FTL, bw=1.5, round_pct=8)
T(sl, "Orchestration decisions are distributed across participating nodes.",
  MAR + Inches(0.3), dist_y + Inches(0.18), CW - Inches(0.6), Inches(0.42),
  sz=FC, bold=True, color=FTL, align=PP_ALIGN.CENTER)
T(sl, "No single global scheduler — each node participates in placement collectively.",
  MAR + Inches(0.3), dist_y + Inches(0.65), CW - Inches(0.6), Inches(0.42),
  sz=FB, bold=False, color=GRY, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 18: Flocky Results ───────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Flocky  ·  Reported Evaluation Results",
          "Key Findings from the Paper", FTL)
# Two big metric callouts
m1w = (CW - Inches(0.4)) / 2
metric(sl, MAR, CT + Inches(0.2), m1w,
       ">97%",
       "of viable neighbours discovered\nin two discovery rounds",
       source="as reported by the paper", accent=FTL)
metric(sl, MAR + m1w + Inches(0.4), CT + Inches(0.2), m1w,
       "84%",
       "less memory than centralized\nKubernetes (reported evaluation)",
       source="as reported by the paper", accent=FTL)
# Divider
R(sl, MAR + m1w + Inches(0.18), CT + Inches(0.3), Pt(1.5), Inches(2.9),
  fill=BDR, border=None)
# Architectural contribution
R(sl, MAR, CT + Inches(3.3), CW, CH - Inches(3.3), fill=SRF, border=FTL, bw=1.5, round_pct=8)
T(sl, "Architectural Contribution",
  MAR + Inches(0.3), CT + Inches(3.5), CW - Inches(0.6), Inches(0.42),
  sz=FS, bold=True, color=FTL)
T(sl, "Flocky demonstrates decentralized edge orchestration without a global controller.\nIntent + node metadata + local discovery enables informed placement at scale.",
  MAR + Inches(0.3), CT + Inches(3.95), CW - Inches(0.6), Inches(0.9),
  sz=FC, bold=False, color=INK, align=PP_ALIGN.CENTER)
add_footer(sl, n, source="As reported by Goethals et al., IEEE TSC 2026")

# ── Slide 19: DATOR Section Divider ───────────────────────────────────────
sl = section_divider(prs, "03", "DATOR",
                     "Self-Organized Neural Network Inference in Dynamic Edges", DOR); n += 1
add_footer(sl, n)

# ── Slide 20: DATOR Paper Info ─────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Paper 3 of 3  ·  DATOR",
          "Self-Organized Neural Network Inference in Dynamic Edge Networks", DOR)
card(sl, MAR, CT, Inches(7.5), CH,
     "Paper Metadata",
     ["Authors: M. Schrauth, M. Thome,",
      "   T. Ohlenforst, F. Kreyß",
      "Venue: Applied Sciences",
      "Volume 15, Issue 23, Article 12615, 2025",
      "DOI: 10.3390/app152312615"],
     style='light', accent=DOR)
rx11 = MAR + Inches(7.7)
rw11 = CW - Inches(7.7)
R(sl, rx11, CT, rw11, CH, fill=DOR, border=None, round_pct=8)
T(sl, "ROLE IN\nSEMINAR",
  rx11 + Inches(0.2), CT + Inches(0.2), rw11 - Inches(0.4), Inches(0.9),
  sz=FS, bold=True, color=WHT)
T(sl, "LAYER 3",
  rx11 + Inches(0.2), CT + Inches(1.2), rw11 - Inches(0.4), Inches(0.5),
  sz=FC, bold=True, color=WHT)
T(sl, "Dynamic\nExecution &\nResilience",
  rx11 + Inches(0.2), CT + Inches(1.75), rw11 - Inches(0.4), Inches(2.1),
  sz=Pt(24), bold=True, color=WHT)
add_footer(sl, n)
SN(sl, "DATOR: self-organized distributed inference in dynamic edge networks with rerouting.")

# ── Slide 21: DATOR Problem ────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "DATOR  ·  Problem",
          "Static Placement Cannot Survive Dynamic Edges", DOR)
chal_d = [
    ("Node Dropout",          ["Nodes disconnect unexpectedly", "Inference pipeline breaks"]),
    ("Intermittent Connectivity", ["Links flicker; packets lost", "Transmission times spike"]),
    ("Heterogeneous Hardware",["Nodes vary in compute power", "One model does not fit all"]),
    ("Changing Workload",     ["Inference demand shifts", "Fixed allocation under/over-serves"]),
]
dw12 = (CW - Inches(0.2)) / 2
dh12 = (CH - Inches(1.2) - Inches(0.1)) / 2
dpos = [
    (MAR, CT),
    (MAR + dw12 + Inches(0.2), CT),
    (MAR, CT + dh12 + Inches(0.1)),
    (MAR + dw12 + Inches(0.2), CT + dh12 + Inches(0.1)),
]
for (dhead, dbody), (dx12, dy12) in zip(chal_d, dpos):
    card(sl, dx12, dy12, dw12, dh12, dhead, dbody, style='light', accent=DOR)
# Design goal
goal_y = CT + 2*dh12 + Inches(0.2)
goal_h = CH - 2*dh12 - Inches(0.2)
R(sl, MAR, goal_y, CW, goal_h, fill=DOR, border=None, round_pct=8)
T(sl, "Design Goal: maintain distributed inference without a central controller or single point of failure.",
  MAR + Inches(0.3), goal_y + goal_h//3, CW - Inches(0.6), goal_h - Inches(0.1),
  sz=FC, bold=True, color=WHT, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 22: DATOR Multi-stage Inference ─────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "DATOR  ·  Multi-Stage Inference",
          "Splitting the Neural Network Across Edge Nodes", DOR)
# Pipeline: INPUT → STAGE1 → STAGE2 → STAGE3 → OUTPUT
# Large nodes: min 1.8 in wide × 1.2 in tall
stage_items = [
    ("INPUT",    None,     "normal",  None),
    ("STAGE 1",  "Node A", "active",  DOR),
    ("STAGE 2",  "Node B", "active",  DOR),
    ("STAGE 3",  "Node C", "active",  DOR),
    ("OUTPUT",   None,     "normal",  None),
]
snw = Inches(1.95)
snh = Inches(1.35)
sarrow = Inches(0.65)
total_w = len(stage_items)*snw + (len(stage_items)-1)*sarrow
sx_start = MAR + (CW - total_w) // 2
ssx = sx_start
ssy = CT + Inches(0.3)
sshapes2 = []
for (stitle, ssub, sstate, sacc) in stage_items:
    node(sl, ssx, ssy, snw, snh, stitle, sub=ssub, state=sstate, accent=DOR)
    sshapes2.append(type('S',(),{'left':ssx,'top':ssy,'width':snw,'height':snh})())
    ssx += snw + sarrow
for i in range(len(sshapes2)-1):
    sa = sshapes2[i]; sb = sshapes2[i+1]
    # Arrow from right edge of sa to left edge of sb
    Arrow(sl, sa.left+sa.width, sa.top+sa.height//2,
          sb.left, sb.top+sb.height//2, color=DOR, width=2.5)
    # "activation" label above arrow midpoint
    mx_a = sa.left + sa.width + sarrow//2
    T(sl, "activation", mx_a - Inches(0.45), ssy - Inches(0.45), Inches(0.9), Inches(0.35),
      sz=FS, bold=False, color=GRY, align=PP_ALIGN.CENTER)
# How it works
how_y = ssy + snh + Inches(0.4)
card(sl, MAR, how_y, CW, CH - snh - Inches(0.75),
     "How It Works",
     ["The DNN is partitioned into layer groups (stages).",
      "Each stage maps to a suitable edge node.",
      "Intermediate activations flow between nodes.",
      "If a node fails — the system finds an alternate path."],
     style='light', accent=DOR)
add_footer(sl, n)

# ── Slide 23: DATOR Local Negotiation ─────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "DATOR  ·  Local Orchestration",
          "Work Assignment Without Central Control", DOR)
# Negotiation flow (left)
neg_items = [
    ("New Work Package",       "Inference stage needs assignment"),
    ("Broadcast Request",      "Current node announces task to neighbours"),
    ("Eligible Nodes Respond", "Nodes with capacity + module reply"),
    ("Node Accepts",           "Best candidate dispatched"),
    ("Stage Executes",         "Result forwarded to next stage"),
]
nfw = Inches(5.8); nfh = Inches(0.78); nfgap = Inches(0.1)
nfy = CT; nfshapes = []
for (nfname, nfdesc) in neg_items:
    node(sl, MAR, nfy, nfw, nfh, nfname, sub=nfdesc, state='normal', accent=DOR)
    nfshapes.append(type('S',(),{'left':MAR,'top':nfy,'width':nfw,'height':nfh})())
    nfy += nfh + nfgap
for i in range(len(nfshapes)-1):
    sa = nfshapes[i]; sb = nfshapes[i+1]
    Arrow(sl, sa.left+sa.width//2, sa.top+sa.height,
          sb.left+sb.width//2, sb.top, color=DOR, width=2)
# Right: decision criteria + Contract Net
rx13 = MAR + nfw + Inches(0.45)
rw13 = CW - nfw - Inches(0.45)
T(sl, "NODE DECISION CRITERIA", rx13, CT, rw13, Inches(0.38),
  sz=FS, bold=True, color=DOR)
crit_items = [
    ("Module present?",      "Does node hold the needed layer group?"),
    ("Utilization OK?",      "Load within acceptance threshold?"),
    ("Hardware capable?",    "Can execute stage efficiently?"),
]
crit_y = CT + Inches(0.45)
for (cq, ca) in crit_items:
    R(sl, rx13, crit_y, rw13, Inches(0.85), fill=SRF, border=BDR, bw=1, round_pct=6)
    T(sl, cq, rx13 + Inches(0.15), crit_y + Inches(0.05),
      rw13 - Inches(0.3), Inches(0.35), sz=FD, bold=True, color=DOR)
    T(sl, ca, rx13 + Inches(0.15), crit_y + Inches(0.38),
      rw13 - Inches(0.3), Inches(0.35), sz=FS, bold=False, color=GRY)
    crit_y += Inches(0.95)
R(sl, rx13, crit_y + Inches(0.1), rw13, Pt(1.5), fill=BDR, border=None)
T(sl, "Weighted-auction factors:",
  rx13, crit_y + Inches(0.25), rw13, Inches(0.38),
  sz=FB, bold=True, color=INK)
T(sl, "Capability · Latency · Energy · Utilization\n(Contract Net Protocol concepts)",
  rx13, crit_y + Inches(0.65), rw13, Inches(0.8),
  sz=FB, bold=False, color=GRY)
add_footer(sl, n)
SN(sl, "DATOR uses broadcast-and-accept negotiation. Weighted auction can consider capability, latency, energy.")

# ── Slide 24: DATOR Resilience ─────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "DATOR  ·  Dynamic Rerouting",
          "Inference Continues Despite Node Failures", DOR)
# Three rows: Normal / Failure / Recovery
# Node dimensions: 1.9 × 1.05 in (>= 1.8 × 1.0 spec)
NW = Inches(1.9); NH = Inches(1.05)
ARR_W = Inches(1.2)   # arrow gap
# Total: label(1.5) + A(1.9) + arrow(1.2) + B(1.9) + arrow(1.2) + C(1.9) = 9.6in
# Horizontal start: after label
LBL_W = Inches(1.55)
A_x = MAR + LBL_W
B_x = A_x + NW + ARR_W
C_x = B_x + NW + ARR_W
D_x = B_x   # D replaces B in recovery row

row_h  = (CH - Inches(0.2)) / 3
row_ys = [CT, CT + row_h + Inches(0.1), CT + 2*row_h + Inches(0.2)]
row_ny = [ry + (row_h - NH) / 2 for ry in row_ys]

row_labels = ["Normal\nA → B → C", "Failure\nB unavailable", "Recovery\nA → D → C"]
row_colors = [FTL, RED, DOR]

# Row labels
for (rl, rc, ry) in zip(row_labels, row_colors, row_ys):
    T(sl, rl, MAR, ry, LBL_W - Inches(0.1), row_h,
      sz=FD, bold=True, color=rc)

# ROW 1: Normal A→B→C (all active/teal)
def draw_node_rect(sl, x, y, w, h, label, state, acc):
    node(sl, x, y, w, h, label, state=state, accent=acc)
    return type('S',(),{'left':x,'top':y,'width':w,'height':h})()

r1a = draw_node_rect(sl, A_x, row_ny[0], NW, NH, "A", 'active', FTL)
r1b = draw_node_rect(sl, B_x, row_ny[0], NW, NH, "B", 'active', FTL)
r1c = draw_node_rect(sl, C_x, row_ny[0], NW, NH, "C", 'active', FTL)
Arrow(sl, r1a.left+r1a.width, r1a.top+r1a.height//2,
      r1b.left, r1b.top+r1b.height//2, color=FTL, width=2.5)
Arrow(sl, r1b.left+r1b.width, r1b.top+r1b.height//2,
      r1c.left, r1c.top+r1c.height//2, color=FTL, width=2.5)

# ROW 2: Failure A→✕→C (B failed/red)
r2a = draw_node_rect(sl, A_x, row_ny[1], NW, NH, "A", 'normal', None)
r2b = draw_node_rect(sl, B_x, row_ny[1], NW, NH, "B  ✕", 'failed', None)
r2c = draw_node_rect(sl, C_x, row_ny[1], NW, NH, "C", 'normal', None)
Arrow(sl, r2a.left+r2a.width, r2a.top+r2a.height//2,
      r2b.left, r2b.top+r2b.height//2, color=RED, width=2.5, dashed=True)
# No arrow A→C directly shown; B→C broken
Conn(sl, r2b.left+r2b.width, r2b.top+r2b.height//2,
     r2c.left, r2c.top+r2c.height//2, color=RED, width=1.5, dashed=True)

# ROW 3: Recovery A→D→C (D in DOR orange)
r3a = draw_node_rect(sl, A_x, row_ny[2], NW, NH, "A", 'normal', None)
r3d = draw_node_rect(sl, D_x, row_ny[2], NW, NH, "D", 'active', DOR)
r3c = draw_node_rect(sl, C_x, row_ny[2], NW, NH, "C", 'normal', None)
Arrow(sl, r3a.left+r3a.width, r3a.top+r3a.height//2,
      r3d.left, r3d.top+r3d.height//2, color=DOR, width=2.5)
Arrow(sl, r3d.left+r3d.width, r3d.top+r3d.height//2,
      r3c.left, r3c.top+r3c.height//2, color=DOR, width=2.5)

# Principle
T(sl, "Fault handling is integrated into the inference path — no centralized recovery controller.",
  MAR, H - Inches(0.52), CW - Inches(0.8), Inches(0.35),
  sz=FB, bold=True, color=DOR)
add_footer(sl, n)

# ── Slide 25: DATOR Evaluation ─────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "DATOR  ·  Evaluation",
          "Performance Across Network Conditions", DOR)
cond_cards = [
    ("Stable",           DOR, ["No node failures", "Baseline performance"]),
    ("Unstable",         DOR, ["Intermittent node failures", "System continues via rerouting"]),
    ("Impaired Comm.",   DOR, ["5% message loss", "Reduced throughput, stable ops"]),
    ("Severe Comm.",     RED, ["~10× longer transmissions", "System approaches saturation"]),
]
ccw = (CW - Inches(0.6)) / 4
cch = Inches(2.5)
ccx = MAR
for (cname, cacc, cbody) in cond_cards:
    card(sl, ccx, CT, ccw, cch, cname, cbody, style='light', accent=cacc)
    ccx += ccw + Inches(0.2)
# Performance pattern
pp_y = CT + cch + Inches(0.2)
pp_h = CH - cch - Inches(0.2)
R(sl, MAR, pp_y, CW, pp_h, fill=SRF, border=DOR, bw=1.5, round_pct=8)
T(sl, "General Performance Pattern",
  MAR + Inches(0.3), pp_y + Inches(0.15), CW - Inches(0.6), Inches(0.38),
  sz=FS, bold=True, color=DOR)
perf_rows = [
    ("Low load",      "Latency remains relatively low",         FTL),
    ("Increasing load", "More nodes join distributed execution", DOR),
    ("Near saturation", "Throughput & latency approach capacity", RED),
]
pry = pp_y + Inches(0.6)
for (pl, pd, pc) in perf_rows:
    R(sl, MAR + Inches(0.3), pry, Inches(0.06), Inches(0.45), fill=pc, border=None)
    T(sl, pl, MAR + Inches(0.5), pry + Inches(0.0), Inches(2.5), Inches(0.45),
      sz=FB, bold=True, color=pc)
    T(sl, pd, MAR + Inches(3.2), pry + Inches(0.0), Inches(8.5), Inches(0.45),
      sz=FB, bold=False, color=GRY)
    pry += Inches(0.7)
T(sl, "Designed to maintain operation even when not all nodes and links are permanently available.",
  MAR + Inches(0.3), pry + Inches(0.1), CW - Inches(0.6), Inches(0.42),
  sz=FB, bold=True, color=INK)
add_footer(sl, n)

# ── Slide 26: Comparison Part 1 ────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Comparative Analysis  ·  1 of 2",
          "Main Question & Method")
# Column setup
row_label_w = Inches(2.6)
col_w26 = (CW - row_label_w - Inches(0.3)) / 3
gap26 = Inches(0.1)
hdr_y = CT; hdr_h = Inches(0.62)
col_xs = [MAR + row_label_w + Inches(0.1),
          MAR + row_label_w + Inches(0.1) + col_w26 + gap26,
          MAR + row_label_w + Inches(0.1) + 2*(col_w26 + gap26)]
hdr_data = [("ASET", ABL), ("Flocky", FTL), ("DATOR", DOR)]
for (hname, hcol), hx in zip(hdr_data, col_xs):
    R(sl, hx, hdr_y, col_w26, hdr_h, fill=hcol, border=None, round_pct=6)
    T(sl, hname, hx, hdr_y, col_w26, hdr_h,
      sz=FC, bold=True, color=WHT, align=PP_ALIGN.CENTER)
# Rows
rows26 = [
    ("Main\nQuestion",
     "Which scheduling policy\nfor the current state?",
     "How can orchestration\nbe decentralized?",
     "How can inference survive\ndynamic nodes & links?"),
    ("Method",
     "Deep Q-Network\nPolicy selection via DRL",
     "Intent + OAM + Metadata\nDecentralized discovery",
     "Multi-stage distributed\nLocal negotiation + rerouting"),
    ("Decision\nLevel",
     "Scheduling policy\nselection",
     "App/component\nplacement",
     "Inference-stage\nassignment & routing"),
]
row_h26 = (CH - hdr_h - Inches(0.1)) / len(rows26)
ry26 = hdr_y + hdr_h + Inches(0.1)
rcolors = [None, ABL, FTL, DOR]
for row in rows26:
    # Row label
    R(sl, MAR, ry26, row_label_w - Inches(0.1), row_h26 - Inches(0.05),
      fill=SRF, border=BDR, bw=0.75, round_pct=4)
    T(sl, row[0], MAR + Inches(0.15), ry26 + Inches(0.12),
      row_label_w - Inches(0.35), row_h26 - Inches(0.2),
      sz=FD, bold=True, color=INK)
    for i, (cell, cx26) in enumerate(zip(row[1:], col_xs)):
        R(sl, cx26, ry26, col_w26, row_h26 - Inches(0.05),
          fill=BG, border=BDR, bw=0.75, round_pct=4)
        T(sl, cell, cx26 + Inches(0.12), ry26 + Inches(0.1),
          col_w26 - Inches(0.24), row_h26 - Inches(0.2),
          sz=FB, bold=False, color=INK, wrap=True)
    ry26 += row_h26
add_footer(sl, n)

# ── Slide 27: Comparison Part 2 ────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Comparative Analysis  ·  2 of 2",
          "Strengths & Challenges")
row_label_w2 = Inches(2.6)
col_w27 = (CW - row_label_w2 - Inches(0.3)) / 3
gap27 = Inches(0.1)
col_xs2 = [MAR + row_label_w2 + Inches(0.1),
           MAR + row_label_w2 + Inches(0.1) + col_w27 + gap27,
           MAR + row_label_w2 + Inches(0.1) + 2*(col_w27 + gap27)]
hdr_y2 = CT; hdr_h2 = Inches(0.62)
hdr_data2 = [("ASET", ABL), ("Flocky", FTL), ("DATOR", DOR)]
for (hname, hcol), hx in zip(hdr_data2, col_xs2):
    R(sl, hx, hdr_y2, col_w27, hdr_h2, fill=hcol, border=None, round_pct=6)
    T(sl, hname, hx, hdr_y2, col_w27, hdr_h2,
      sz=FC, bold=True, color=WHT, align=PP_ALIGN.CENTER)
rows27 = [
    ("Key\nStrength",
     "Workload / network\nadaptation",
     "Decentralized coordination\n& local knowledge",
     "Self-organization\n& resilience"),
    ("Challenge\nAddressed",
     "Changing system\nconditions",
     "Scalable decentralized\norchestration",
     "Node / link failures\n& dynamic environments"),
    ("Not\nCompeting",
     "Adaptive decision\nmaking (layer 1)",
     "Decentralized coordination\n(layer 2)",
     "Dynamic execution\n& recovery (layer 3)"),
]
row_h27 = (CH - hdr_h2 - Inches(0.1)) / len(rows27)
ry27 = hdr_y2 + hdr_h2 + Inches(0.1)
for row in rows27:
    R(sl, MAR, ry27, row_label_w2 - Inches(0.1), row_h27 - Inches(0.05),
      fill=SRF, border=BDR, bw=0.75, round_pct=4)
    T(sl, row[0], MAR + Inches(0.15), ry27 + Inches(0.12),
      row_label_w2 - Inches(0.35), row_h27 - Inches(0.2),
      sz=FD, bold=True, color=INK)
    for cell, cx27 in zip(row[1:], col_xs2):
        R(sl, cx27, ry27, col_w27, row_h27 - Inches(0.05),
          fill=BG, border=BDR, bw=0.75, round_pct=4)
        T(sl, cell, cx27 + Inches(0.12), ry27 + Inches(0.1),
          col_w27 - Inches(0.24), row_h27 - Inches(0.2),
          sz=FB, bold=False, color=INK, wrap=True)
    ry27 += row_h27
add_footer(sl, n)

# ── Slide 28: Synthesis ─────────────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Synthesis", "Three Papers · One Problem Space")
# 3 paper boxes
syn_cards = [
    ("ASET",   ABL,  "Adaptive Scheduling",         "State-aware policy selection via DRL"),
    ("Flocky", FTL,  "Decentralized Orchestration",  "Intent + OAM + local discovery"),
    ("DATOR",  DOR,  "Dynamic Resilient Inference",   "Multi-stage + rerouting + self-org"),
]
scw = (CW - Inches(0.4)) / 3
sch = Inches(2.0)
scx = MAR
sshapes3 = []
for (sname, sacc, stheme, sdesc) in syn_cards:
    R(sl, scx, CT, scw, sch, fill=SRF, border=BDR, bw=1.5, round_pct=8)
    R(sl, scx, CT, scw, Inches(0.06), fill=sacc, border=None)
    T(sl, sname, scx + Inches(0.2), CT + Inches(0.12), scw - Inches(0.4), Inches(0.6),
      sz=Pt(28), bold=True, color=sacc)
    T(sl, stheme, scx + Inches(0.2), CT + Inches(0.72), scw - Inches(0.4), Inches(0.38),
      sz=FB, bold=True, color=INK)
    T(sl, sdesc, scx + Inches(0.2), CT + Inches(1.12), scw - Inches(0.4), Inches(0.68),
      sz=FB, bold=False, color=GRY)
    # Make fake shape for arrow computation
    mid = scx + scw // 2
    sshapes3.append(type('S',(),{
        'left':int(mid - Inches(0.05)), 'top':int(CT+sch),
        'width':Inches(0.1), 'height':Inches(0.1)
    })())
    scx += scw + Inches(0.2)
# Converging arrows
syn_box_y = CT + sch + Inches(0.5)
syn_box_mid_x = MAR + CW // 2
for ss in sshapes3:
    Arrow(sl, ss.left, ss.top, syn_box_mid_x, syn_box_y, color=GRY, width=1.5)
# Combined box
syn_h = CH - sch - Inches(0.55)
R(sl, MAR, syn_box_y, CW, syn_h, fill=INK, border=None, round_pct=8)
T(sl, "INTELLIGENT DISTRIBUTED EDGE AI",
  MAR + Inches(0.3), syn_box_y + Inches(0.25), CW - Inches(0.6), Inches(0.6),
  sz=Pt(26), bold=True, color=WHT, align=PP_ALIGN.CENTER)
# Lifecycle
lifecycle = ["Observe", "Decide", "Coordinate", "Execute", "Adapt", "Recover"]
lw = (CW - Inches(0.6)) / len(lifecycle)
lx_ = MAR + Inches(0.3)
lcolors2 = [ABL, ABL, FTL, FTL, DOR, DOR]
for (ltxt, lcol) in zip(lifecycle, lcolors2):
    R(sl, lx_, syn_box_y + Inches(0.95), lw - Inches(0.08), Inches(0.62),
      fill=None, border=lcol, bw=1.5, round_pct=6)
    T(sl, ltxt, lx_, syn_box_y + Inches(0.95), lw - Inches(0.08), Inches(0.62),
      sz=FD, bold=True, color=lcol, align=PP_ALIGN.CENTER)
    lx_ += lw
T(sl, "Adaptive scheduling  +  Decentralized orchestration  +  Resilient inference",
  MAR + Inches(0.3), syn_box_y + Inches(1.65), CW - Inches(0.6), Inches(0.5),
  sz=FB, bold=False, color=GRY, align=PP_ALIGN.CENTER)
add_footer(sl, n)

# ── Slide 29: Future Directions ────────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Future Directions", "Open Questions in Edge AI Coordination")
# Left: what the papers show
left_w = Inches(5.5)
card(sl, MAR, CT, left_w, Inches(2.4),
     "Demonstrated by the Three Papers",
     ["Adaptive scheduling (ASET)",
      "Decentralized orchestration (Flocky)",
      "Dynamic self-organization (DATOR)"],
     style='light')
# Right + bottom: future items
future_items = [
    ("Richer state-aware scheduling",    None),
    ("Decentralized discovery at scale", None),
    ("Energy-aware dynamic routing",     None),
    ("Multimodal AI inference at edge",  None),
    ("Cross-layer integration",          None),
    ("Fault tolerance in large clusters",None),
]
fw29 = (CW - left_w - Inches(0.4)) / 2
fh29 = (Inches(2.4) - Inches(0.1)) / 3
fx29 = MAR + left_w + Inches(0.4)
fy29 = CT
for i, (fname, _) in enumerate(future_items):
    col29 = i % 2
    row29 = i // 2
    fxi = fx29 + col29 * (fw29 + Inches(0.15))
    fyi = fy29 + row29 * (fh29 + Inches(0.1))
    R(sl, fxi, fyi, fw29, fh29, fill=SRF, border=BDR, bw=1, round_pct=6)
    T(sl, fname, fxi + Inches(0.15), fyi + Inches(0.05),
      fw29 - Inches(0.3), fh29 - Inches(0.1),
      sz=FB, bold=False, color=INK, wrap=True)
# Integration vision box
iv_y = CT + Inches(2.55)
card(sl, MAR, iv_y, CW, CH - Inches(2.55),
     "Integration Vision",
     ["A unified edge system that can simultaneously: observe state,",
      "select resources, coordinate nodes, adapt, and recover."],
     style='light')
add_footer(sl, n)

# ── Slide 30: Summary + References ────────────────────────────────────────
sl = prs.slides.add_slide(_blank); n += 1
add_bg(sl)
add_title(sl, "Summary", "Key Takeaways & References")
# 3 summary cards
sum_items = [
    ("ASET",   ABL,
     ["Adaptive scheduling via Deep Q-Learning",
      "Meta-scheduler: selects the policy, not the node"]),
    ("Flocky", FTL,
     ["Decentralized orchestration via intent + OAM",
      "Local discovery and capability matching"]),
    ("DATOR",  DOR,
     ["Self-organized multi-stage distributed inference",
      "Dynamic rerouting; resilience without central control"]),
]
scw2 = (CW - Inches(0.4)) / 3
sch2 = Inches(1.85)
scx2 = MAR
for (sname, sacc, sbody) in sum_items:
    card(sl, scx2, CT, scw2, sch2, sname, sbody, style='light', accent=sacc)
    scx2 += scw2 + Inches(0.2)
# Core equation
eq_y = CT + sch2 + Inches(0.2)
R(sl, MAR, eq_y, CW, Inches(0.55), fill=INK, border=None, round_pct=6)
T(sl, "Adaptive Scheduling  +  Decentralized Orchestration  +  Resilient Inference  =  Intelligent Edge AI",
  MAR + Inches(0.3), eq_y + Inches(0.07), CW - Inches(0.6), Inches(0.42),
  sz=FD, bold=True, color=WHT, align=PP_ALIGN.CENTER)
# References
ref_y = eq_y + Inches(0.72)
ref_h = CB - ref_y
R(sl, MAR, ref_y, CW, ref_h, fill=SRF, border=BDR, bw=1, round_pct=6)
T(sl, "References", MAR + Inches(0.2), ref_y + Inches(0.12),
  CW - Inches(0.4), Inches(0.38), sz=FS, bold=True, color=GRY)
refs30 = [
    "[1]  G. Castellano et al., \"Scheduling Inference Workloads in the Computing Continuum with RL,\" IEEE ICDCS Wkshps, 2025, pp. 297-302. DOI: 10.1109/ICDCSW63273.2025.00056",
    "[2]  T. Goethals et al., \"Flocky: Decentralized Intent-based Edge Orchestration Using OAM,\" IEEE TSC, vol. 19, no. 1, pp. 380-395, 2026. DOI: 10.1109/TSC.2025.3631387",
    "[3]  M. Schrauth et al., \"Self-Organized Neural Network Inference in Dynamic Edge Networks,\" Applied Sciences, vol. 15, no. 23, 12615, 2025. DOI: 10.3390/app152312615",
]
rfy = ref_y + Inches(0.55)
for ref30 in refs30:
    T(sl, ref30, MAR + Inches(0.2), rfy, CW - Inches(0.4), Inches(0.52),
      sz=FS, bold=False, color=GRY, wrap=True)
    rfy += Inches(0.68)
add_footer(sl, n)

# ══════════════════════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════════════════════
OUT = r"c:\Malenia\TASK_MESH\Adaptive_Distributed_AI_Inference_Seminar.pptx"
prs.save(OUT)
print(f"Saved: {OUT}")
print(f"Total slides: {len(prs.slides)}")
